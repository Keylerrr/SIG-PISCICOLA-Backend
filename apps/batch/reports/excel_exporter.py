import io

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, GradientFill, PatternFill, Side
)
from openpyxl.utils import get_column_letter

from .constants import MODULE_COLORS, MODULE_FEEDING, MODULE_BIOMETRY, MODULE_HARVEST, MODULE_LABELS

# ── Paleta (misma que el PDF) ──────────────────────────────────────────────
C_TEAL_DARK  = "0D4F5C"
C_TEAL_MID   = "1A7A8A"
C_TEAL_LIGHT = "2AABB9"
C_ACCENT     = "4ECDC4"
C_GOLD       = "F4A261"
C_BG_LIGHT   = "F0F7F8"
C_BG_ROW     = "E8F4F6"
C_GRAY_TEXT  = "4A5568"
C_GRAY_LIGHT = "CBD5E0"
C_WHITE      = "FFFFFF"

# ── Estilos base ──────────────────────────────────────────────────────────
FONT_BASE   = "Arial"

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)

def _font(bold=False, size=10, color=C_GRAY_TEXT, italic=False) -> Font:
    return Font(name=FONT_BASE, bold=bold, size=size, color=color, italic=italic)

def _border_thin(color=C_GRAY_LIGHT) -> Border:
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def _border_bottom(color=C_TEAL_MID) -> Border:
    return Border(bottom=Side(style="medium", color=color))

def _align(h="left", v="center", wrap=False) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


# ── Helpers de escritura ──────────────────────────────────────────────────
def _autosize_columns(ws, min_w=10, max_w=45):
    for col_idx, col_cells in enumerate(ws.columns, 1):
        width = max(
            (len(str(c.value)) for c in col_cells if c.value is not None),
            default=min_w,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 3, max_w)


def _write_section_title(ws, text: str, hex_color: str = C_TEAL_DARK, ncols: int = 1):
    """Fila de título de sección con fondo de color (igual al PDF)."""
    row = ws.max_row + 1
    ws.row_dimensions[row].height = 20
    cell = ws.cell(row=row, column=1, value=text)
    cell.font      = _font(bold=True, size=12, color=C_WHITE)
    cell.fill      = _fill(hex_color)
    cell.alignment = _align("left", "center")
    # Merge si hay más de una columna
    if ncols > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    return row


def _write_subheading(ws, text: str, ncols: int = 1):
    row = ws.max_row + 1
    ws.row_dimensions[row].height = 16
    cell = ws.cell(row=row, column=1, value=text)
    cell.font      = _font(bold=True, size=10, color=C_TEAL_DARK)
    cell.alignment = _align("left", "center")
    cell.border    = _border_bottom()
    if ncols > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)


def _write_header_row(ws, headers: list[str], hex_color: str = C_TEAL_MID) -> int:
    row = ws.max_row + 1
    ws.row_dimensions[row].height = 15
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font      = _font(bold=True, size=9, color=C_WHITE)
        cell.fill      = _fill(hex_color)
        cell.alignment = _align("center", "center", wrap=True)
        cell.border    = _border_thin(C_WHITE)
    return row


def _write_data_rows(ws, rows: list[list], start_col: int = 1):
    """Escribe filas de datos con alternancia de color."""
    for i, row_data in enumerate(rows):
        row = ws.max_row + 1
        ws.row_dimensions[row].height = 13
        bg = C_BG_ROW if i % 2 == 0 else C_WHITE
        for col, value in enumerate(row_data, start_col):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font      = _font(size=9)
            cell.fill      = _fill(bg)
            cell.alignment = _align("left", "center", wrap=True)
            cell.border    = _border_thin()


def _write_table(ws, headers: list[str], rows: list[list],
                 hex_color: str = C_TEAL_MID):
    """Escribe encabezado + filas de datos en una sola llamada."""
    if not rows:
        r = ws.max_row + 1
        cell = ws.cell(row=r, column=1, value="Sin registros para este subgrupo.")
        cell.font = _font(italic=True, size=9, color=C_GRAY_TEXT)
        return
    _write_header_row(ws, headers, hex_color)
    _write_data_rows(ws, rows)


def _spacer(ws, n: int = 1):
    for _ in range(n):
        ws.append([None])


# ── Hoja de Resumen ───────────────────────────────────────────────────────
def _build_summary_sheet(ws, report_data: dict):
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 40

    batch = report_data["batch"]
    ncols = 6

    # Título principal
    _write_section_title(ws, "Reporte histórico del proceso productivo", C_TEAL_DARK, ncols)

    # Metadata del lote
    _spacer(ws)
    _write_subheading(ws, "Información del lote", ncols)
    meta_rows = [
        ("Lote",              batch.get("code", "—")),
        ("Especie",           batch.get("specie", "—")),
        ("Estado biológico",  batch.get("biological_state", "—")),
        ("Estado",            batch.get("status", "—")),
        ("Cantidad inicial",  batch.get("initial_quantity", "—")),
        ("Granja",            batch.get("farm_name", "—")),
        ("Generado",          report_data.get("generated_at", "—")),
        ("Usuario",           report_data.get("generated_by", "—")),
    ]
    if report_data.get("cycle_id"):
        meta_rows.append(("Ciclo consultado (ID)", report_data["cycle_id"]))
    for label, value in meta_rows:
        row = ws.max_row + 1
        lc = ws.cell(row=row, column=1, value=label)
        vc = ws.cell(row=row, column=2, value=value)
        lc.font      = _font(bold=True, size=9, color=C_TEAL_DARK)
        vc.font      = _font(size=9)
        lc.alignment = _align()
        vc.alignment = _align()
        lc.border    = _border_thin(C_GRAY_LIGHT)
        vc.border    = _border_thin(C_GRAY_LIGHT)

    # Indicadores de productividad
    _spacer(ws)
    _write_section_title(ws, "Indicadores de productividad", C_TEAL_DARK, ncols)
    _spacer(ws)
    ind_headers = ["Indicador", "Valor", "Estado", "Clasificación", "Recomendación", "Notas"]
    ind_rows = [
        [
            i.get("name", ""),
            i.get("value", ""),
            i.get("status", ""),
            i.get("classification", ""),
            i.get("recommendation", ""),
            i.get("notes", ""),
        ]
        for i in report_data.get("productivity_indicators", [])
    ]
    _write_table(ws, ind_headers, ind_rows, C_TEAL_MID)

    # Módulos incluidos
    _spacer(ws)
    _write_section_title(ws, "Módulos incluidos", C_TEAL_DARK, 2)
    _spacer(ws)
    mod_rows = [
        [report_data["modules"][k]["label"], "Con datos" if report_data["modules"][k]["has_data"] else "Sin registros"]
        for k in report_data["selected_modules"]
    ]
    _write_table(ws, ["Módulo", "Estado"], mod_rows, C_TEAL_MID)

    if report_data.get("empty_module_labels"):
        _spacer(ws)
        _write_subheading(ws, "Módulos seleccionados sin información", 2)
        for label in report_data["empty_module_labels"]:
            r = ws.max_row + 1
            cell = ws.cell(row=r, column=1, value=label)
            cell.font = _font(size=9, italic=True, color=C_GRAY_TEXT)

    # Ciclos asociados
    ctx = report_data["context"]
    if ctx.get("cycles"):
        _spacer(ws)
        _write_section_title(ws, "Ciclos asociados", C_TEAL_DARK, 5)
        _spacer(ws)
        cycle_rows = [
            [c["name"], c["pond"], c["state"], c["start_date"], c["finish_date"] or "—"]
            for c in ctx["cycles"]
        ]
        _write_table(ws, ["Nombre", "Estanque", "Estado", "Inicio", "Fin"], cycle_rows, C_TEAL_MID)

    _autosize_columns(ws)


# ── Hojas de módulos ──────────────────────────────────────────────────────
def _filter_by_type(records: list[dict], tipo: str) -> list[dict]:
    return [r for r in records if r.get("tipo") == tipo]


def _build_feeding_sheet(ws, records: list[dict], color: str):
    ws.sheet_view.showGridLines = False

    # Cronogramas
    cronogramas = _filter_by_type(records, "cronograma")
    _write_section_title(ws, "Detalles de cronogramas", color, 13)
    _spacer(ws)
    _write_table(ws,
        ["Nombre", "Etapa", "Forma alimento", "Tamaño pellet (mm)", "% Alimentación",
         "Veces/día", "Int. raciones (min)", "Int. jornadas (días)", "FCA esperado",
         "Ganancia diaria (g)", "Peso mín (g)", "Peso máx (g)", "Comentarios"],
        [[
            r.get("nombre","—"), r.get("etapa","—"), r.get("forma_alimento","—"),
            r.get("tamano_pellet_mm","—"), r.get("porcentaje_alimentacion","—"),
            r.get("veces_por_dia","—"), r.get("intervalo_entre_raciones_min","—"),
            r.get("intervalo_entre_jornadas_dias","—"), r.get("esperado_fca","—"),
            r.get("ganancia_diaria_g","—"), r.get("peso_min_aceptable_g","—"),
            r.get("peso_max_aceptable_g","—"), r.get("comentarios",""),
        ] for r in cronogramas],
        color,
    )

    # Cronogramas de alimentación (programas únicos)
    plans = _filter_by_type(records, "plan")
    _spacer(ws)
    _write_section_title(ws, "Cronogramas de alimentación", color, 4)
    _spacer(ws)
    seen, prog_rows = set(), []
    for p in plans:
        name = p.get("programa", "—")
        if name not in seen:
            seen.add(name)
            prog_rows.append([name, p.get("ciclo","—"), p.get("inicio","—"), p.get("fin","—")])
    _write_table(ws, ["Programa", "Ciclo", "Inicio", "Fin"], prog_rows, color)

    # Planes de alimentación
    _spacer(ws)
    _write_section_title(ws, "Planes de alimentación", color, 4)
    _spacer(ws)
    _write_table(ws, ["Ciclo", "Programa", "Inicio", "Fin"],
        [[p.get("ciclo","—"), p.get("programa","—"), p.get("inicio","—"), p.get("fin","—")] for p in plans],
        color,
    )

    # Eventos
    events = _filter_by_type(records, "evento")
    _spacer(ws)
    _write_section_title(ws, "Eventos de alimentación", color, 7)
    _spacer(ws)
    _write_table(ws, ["Ciclo", "Fecha", "Hora", "Ración", "Estado", "Cantidad", "Unidad"],
        [[e.get("ciclo","—"), e.get("fecha","—"), e.get("hora","—"), e.get("racion","—"),
          e.get("estado","—"), e.get("cantidad","—"), e.get("unidad","—")] for e in events],
        color,
    )
    _autosize_columns(ws)


def _build_biometry_sheet(ws, records: list[dict], color: str):
    ws.sheet_view.showGridLines = False

    controls = _filter_by_type(records, "control")
    _write_section_title(ws, "Controles realizados", color, 9)
    _spacer(ws)
    _write_table(ws,
        ["Ciclo", "Estanque", "Fecha", "Muestra", "Vivos",
         "Peso prom (g)", "Mortalidad (%)", "Biomasa (kg)", "FCA"],
        [[c.get("ciclo","—"), c.get("estanque","—"), c.get("fecha","—"),
          c.get("muestra","—"), c.get("vivos","—"), c.get("peso_prom_g","—"),
          c.get("mortalidad_%","—"), c.get("biomasa_kg","—"), c.get("fca","—")] for c in controls],
        color,
    )

    evaluations = _filter_by_type(records, "evaluacion")
    _spacer(ws)
    _write_section_title(ws, "Evaluaciones realizadas", color, 6)
    _spacer(ws)
    _write_table(ws,
        ["Ciclo", "Estanque", "Fecha", "Muestra", "Peso prom (g)", "Mortalidad"],
        [[e.get("ciclo","—"), e.get("estanque","—"), e.get("fecha","—"),
          e.get("muestra","—"), e.get("peso_prom_g","—"), e.get("mortalidad","—")] for e in evaluations],
        color,
    )
    _autosize_columns(ws)


def _build_harvest_sheet(ws, records: list[dict], color: str):
    ws.sheet_view.showGridLines = False

    harvests = _filter_by_type(records, "cosecha")
    _write_section_title(ws, "Cosechas", color, 5)
    _spacer(ws)
    _write_table(ws, ["Ciclo", "Fecha", "Tipo cosecha", "Peces", "Peso total (g)"],
        [[h.get("ciclo","—"), h.get("fecha","—"), h.get("tipo_cosecha","—"),
          h.get("peces","—"), h.get("peso_total_g","—")] for h in harvests],
        color,
    )

    sources = _filter_by_type(records, "fuente_cosecha")
    _spacer(ws)
    _write_section_title(ws, "Fuentes de cosecha", color, 4)
    _spacer(ws)
    _write_table(ws, ["Cosecha ID", "Peces", "Peso (g)", "Trazabilidad"],
        [[s.get("cosecha_id","—"), s.get("peces","—"), s.get("peso_g","—"), s.get("trazabilidad","—")]
         for s in sources],
        color,
    )

    clas = _filter_by_type(records, "clasificacion")
    _spacer(ws)
    _write_section_title(ws, "Clasificaciones", color, 4)
    _spacer(ws)
    _write_table(ws, ["Cosecha ID", "Categoría", "Peces", "Peso (g)"],
        [[c.get("cosecha_id","—"), c.get("categoria","—"), c.get("peces","—"), c.get("peso_g","—")]
         for c in clas],
        color,
    )

    deriv = _filter_by_type(records, "derivacion_lote")
    _spacer(ws)
    _write_section_title(ws, "Derivaciones de lote", color, 3)
    _spacer(ws)
    _write_table(ws, ["Clasificación ID", "Peces", "Peso (g)"],
        [[d.get("clasificacion_id","—"), d.get("peces","—"), d.get("peso_g","—")] for d in deriv],
        color,
    )
    _autosize_columns(ws)


def _build_generic_sheet(ws, records: list[dict], color: str, label: str):
    ws.sheet_view.showGridLines = False
    if not records:
        r = ws.max_row + 1
        cell = ws.cell(row=r, column=1,
            value="No se encontró información registrada para este módulo en el alcance consultado.")
        cell.font = _font(italic=True, size=9, color=C_GRAY_TEXT)
        return

    keys = []
    for rec in records:
        for k in rec:
            if k not in keys:
                keys.append(k)

    _write_section_title(ws, label, color, len(keys))
    _spacer(ws)
    headers = [k.replace("_", " ").title() for k in keys]
    rows    = [[rec.get(k, "") for k in keys] for rec in records]
    _write_table(ws, headers, rows, color)
    _autosize_columns(ws)


# ── Exportador principal ──────────────────────────────────────────────────
def export_production_report_xlsx(report_data: dict) -> bytes:
    wb = Workbook()
    _build_summary_sheet(wb.active, report_data)

    for module_key in report_data["selected_modules"]:
        mod        = report_data["modules"][module_key]
        color      = (MODULE_COLORS.get(module_key) or C_TEAL_DARK).lstrip("#")
        sheet_name = mod["label"][:31]
        ws         = wb.create_sheet(title=sheet_name)

        if module_key == MODULE_FEEDING:
            _build_feeding_sheet(ws, mod["records"], color)
        elif module_key == MODULE_BIOMETRY:
            _build_biometry_sheet(ws, mod["records"], color)
        elif module_key == MODULE_HARVEST:
            _build_harvest_sheet(ws, mod["records"], color)
        else:
            _build_generic_sheet(ws, mod["records"], color, mod["label"])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()